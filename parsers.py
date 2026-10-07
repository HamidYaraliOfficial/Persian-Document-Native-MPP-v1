from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from .models import Dependency, ProjectSpec, TaskSpec
from .text_utils import normalize_text, parse_dependencies, parse_duration_days, parse_iso_or_local_date, extract_field

SUPPORTED = {".pdf", ".doc", ".docx"}


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("برای PDF باید pypdf نصب باشد: pip install pypdf") from exc
    chunks: list[str] = []
    with path.open("rb") as fh:
        reader = PdfReader(fh)
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            chunks.append(f"\n[PAGE {i}]\n{text}")
    return "\n".join(chunks)


def _iter_docx_blocks(path: Path) -> Iterable[tuple[str, object]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("برای DOCX باید python-docx نصب باشد: pip install python-docx") from exc
    doc = Document(path)
    for p in doc.paragraphs:
        if p.text.strip():
            yield "paragraph", p
    for idx, table in enumerate(doc.tables, start=1):
        yield f"table:{idx}", table


def extract_docx_text(path: Path) -> str:
    chunks: list[str] = []
    for kind, obj in _iter_docx_blocks(path):
        if kind == "paragraph":
            chunks.append(obj.text)
        else:
            rows = []
            for row in obj.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            chunks.append("\n".join(rows))
    return "\n".join(chunks)


def _heading_level(paragraph) -> int | None:
    style = getattr(paragraph, "style", None)
    name = (getattr(style, "name", "") or "").lower()
    m = re.search(r"(?:heading|عنوان|تیتر)\s*([1-9])", name)
    if m:
        return int(m.group(1))
    # Persian/Latin numbered headings: 1. / 1.2 / ۱-۲-
    text = normalize_text(paragraph.text)
    m = re.match(r"^((?:\d+[.)/-]?\s*){1,6})\S", text)
    if m and len(re.findall(r"\d+", m.group(1))) <= 6:
        return len(re.findall(r"\d+", m.group(1)))
    return None


def _list_level(paragraph) -> int | None:
    style_name = (getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
    if any(x in style_name for x in ("list bullet", "list number", "bullet", "number")):
        return 2
    text = normalize_text(paragraph.text)
    if re.match(r"^(?:[-*•▪◦]|\d+[.)]|[الفبجده][.)])\s+", text):
        return 2
    return None


def _clean_task_name(name: str) -> str:
    name = normalize_text(name)
    name = re.sub(r"^(?:[-*•▪◦]|\d+[.)]|[الفبجده][.)])\s*", "", name)
    # Remove inline project metadata after the task title.
    name = re.split(r"\s*[—–-]\s*(?=(?:مدت|زمان|duration|شروع|تاریخ شروع|start|پایان|تاریخ پایان|finish|end|پیش(?:‌|\s*)نیاز|پیشنیاز|predecessor|depends\s*on)\s*[:=])", name, maxsplit=1, flags=re.I)[0]
    return name.strip(" :-–—") or "وظیفه بدون عنوان"


def _header_index(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    for i, h in enumerate(headers):
        n = normalize_text(h).lower()
        if any(c in n for c in candidates):
            return i
    return None


def parse_docx(path: Path, project_name: str | None = None) -> ProjectSpec:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX") from exc
    doc = Document(path)
    tasks: list[TaskSpec] = []
    inferred_project = project_name or path.stem
    last_level = 1

    for pidx, p in enumerate(doc.paragraphs, start=1):
        text = normalize_text(p.text)
        if not text:
            continue
        level = _heading_level(p)
        if level is None:
            level = _list_level(p)
        if level is None:
            # Field-rich lines can still represent tasks, e.g. "طراحی API — مدت: 5 روز"
            if re.search(r"(?:مدت|duration|شروع|start|پیش(?:‌|\s*)نیاز|predecessor)\s*[:=]", text, re.I):
                level = max(1, last_level)
            else:
                continue
        last_level = level
        name = _clean_task_name(text)
        duration_txt = extract_field(text, [r"(?:مدت|duration|زمان)\s*[:=]\s*([^,؛;|]+)"])
        start_txt = extract_field(text, [r"(?:شروع|تاریخ شروع|start)\s*[:=]\s*([^,؛;|]+)"])
        finish_txt = extract_field(text, [r"(?:پایان|تاریخ پایان|finish|end)\s*[:=]\s*([^,؛;|]+)"])
        dep_field = extract_field(text, [r"(?:پیش(?:‌|\s*)نیاز|پیشنیاز|predecessors?|depends\s*on|بعد از)\s*[:=]\s*(.+)$"])
        deps = [Dependency(*x) for x in parse_dependencies(dep_field or "")]
        dur = parse_duration_days(duration_txt or "")
        task = TaskSpec(
            name=name,
            level=level,
            duration_days=dur,
            start=parse_iso_or_local_date(start_txt or ""),
            finish=parse_iso_or_local_date(finish_txt or ""),
            predecessors=deps,
            milestone=(dur == 0),
            source_line=pidx,
        )
        tasks.append(task)

    # Tables are generally the most reliable source of structured project data.
    for tindex, table in enumerate(doc.tables, start=1):
        rows = [[normalize_text(cell.text) for cell in row.cells] for row in table.rows]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        headers = rows[0]
        name_i = _header_index(headers, ("task", "name", "activity", "عنوان", "فعالیت", "وظیفه"))
        dur_i = _header_index(headers, ("duration", "مدت", "زمان"))
        start_i = _header_index(headers, ("start", "شروع", "تاریخ شروع"))
        finish_i = _header_index(headers, ("finish", "end", "پایان", "تاریخ پایان"))
        pred_i = _header_index(headers, ("predecessor", "dependency", "depends", "پیش نیاز", "پیش‌نیاز", "پیشنیاز"))
        res_i = _header_index(headers, ("resource", "مسئول", "منبع"))
        note_i = _header_index(headers, ("note", "notes", "شرح", "توضیحات"))
        level_i = _header_index(headers, ("level", "outline", "سطح", "سطح wbs"))
        if name_i is None:
            # Accept simple 2+ column tables where first column is the task name.
            name_i = 0 if len(headers) >= 2 else None
        if name_i is None:
            continue
        for ridx, row in enumerate(rows[1:], start=2):
            if name_i >= len(row) or not row[name_i].strip():
                continue
            def get(i): return row[i] if i is not None and i < len(row) else ""
            dur = parse_duration_days(get(dur_i))
            deps = [Dependency(*x) for x in parse_dependencies(get(pred_i))]
            lvl = int(parse_duration_days(get(level_i)) or 1) if get(level_i) else 1
            task = TaskSpec(
                name=_clean_task_name(get(name_i)), level=max(1, min(20, lvl)),
                duration_days=dur,
                start=parse_iso_or_local_date(get(start_i)),
                finish=parse_iso_or_local_date(get(finish_i)),
                predecessors=deps,
                resource_names=[x.strip() for x in re.split(r"[,;،؛]+", get(res_i)) if x.strip()],
                notes=get(note_i), milestone=(dur == 0), source_table=tindex,
            )
            tasks.append(task)

    start = next((t.start for t in tasks if t.start), None)
    return ProjectSpec(name=inferred_project, start=start, tasks=tasks)


def parse_pdf(path: Path, project_name: str | None = None) -> ProjectSpec:
    text = extract_pdf(path)
    lines = [normalize_text(x) for x in text.splitlines() if normalize_text(x)]
    tasks: list[TaskSpec] = []
    inferred_project = project_name or path.stem
    current_level = 1
    for idx, line in enumerate(lines, start=1):
        if re.fullmatch(r"\[PAGE\s+\d+\]", line, re.I):
            continue
        heading = re.match(r"^((?:\d+[.)/-]?\s*){1,6})(.+)$", line)
        bullet = re.match(r"^(?:[-*•▪◦]|\d+[.)]|[الفبجده][.)])\s+(.+)$", line)
        if heading:
            nums = re.findall(r"\d+", heading.group(1))
            current_level = max(1, len(nums))
            name = _clean_task_name(heading.group(2))
        elif bullet:
            name = _clean_task_name(bullet.group(1))
        elif re.search(r"(?:مدت|duration|شروع|start|پیش(?:‌|\s*)نیاز|predecessor)\s*[:=]", line, re.I):
            name = _clean_task_name(line)
        else:
            continue
        duration_txt = extract_field(line, [r"(?:مدت|duration|زمان)\s*[:=]\s*([^,؛;|]+)"])
        start_txt = extract_field(line, [r"(?:شروع|تاریخ شروع|start)\s*[:=]\s*([^,؛;|]+)"])
        finish_txt = extract_field(line, [r"(?:پایان|تاریخ پایان|finish|end)\s*[:=]\s*([^,؛;|]+)"])
        pred_txt = extract_field(line, [r"(?:پیش(?:‌|\s*)نیاز|پیشنیاز|predecessors?|depends\s*on|بعد از)\s*[:=]\s*(.+)$"])
        dur = parse_duration_days(duration_txt or "")
        tasks.append(TaskSpec(
            name=name, level=current_level, duration_days=dur,
            start=parse_iso_or_local_date(start_txt or ""),
            finish=parse_iso_or_local_date(finish_txt or ""),
            predecessors=[Dependency(*x) for x in parse_dependencies(pred_txt or "")],
            milestone=(dur == 0), source_line=idx,
        ))
    return ProjectSpec(name=inferred_project, start=next((t.start for t in tasks if t.start), None), tasks=tasks)


def _convert_legacy_doc_to_docx(path: Path) -> Path:
    if __import__("platform").system() != "Windows":
        raise RuntimeError("فایل .doc قدیمی فقط روی Windows با Microsoft Word قابل تبدیل خودکار است. آن را به .docx ذخیره کنید.")
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise RuntimeError("برای پشتیبانی از .doc، pywin32 نصب کنید: pip install pywin32") from exc
    import tempfile
    out = Path(tempfile.mkdtemp(prefix="doc_to_docx_")) / f"{path.stem}.docx"
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(str(path.resolve()), ReadOnly=True)
        # wdFormatXMLDocument = 12
        doc.SaveAs2(str(out), FileFormat=12)
        return out
    finally:
        try:
            if doc is not None:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit(False)
        except Exception:
            pass


def parse_file(path: str | Path, project_name: str | None = None) -> ProjectSpec:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    if p.suffix.lower() not in SUPPORTED:
        raise ValueError(f"فرمت پشتیبانی نمی‌شود: {p.suffix}. فرمت‌های مجاز: PDF, DOC, DOCX")
    if p.suffix.lower() == ".docx":
        return parse_docx(p, project_name)
    if p.suffix.lower() == ".doc":
        return parse_docx(_convert_legacy_doc_to_docx(p), project_name)
    return parse_pdf(p, project_name)
